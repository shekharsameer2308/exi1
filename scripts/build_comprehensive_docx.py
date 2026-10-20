import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import re
import os

def sanitize(s):
    if not isinstance(s, str):
        return s
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', s)

doc = docx.Document()

# Page Margins (1 inch)
for section in doc.sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

COLOR_SLATE = RGBColor(0x47, 0x55, 0x69)
COLOR_TEXT = RGBColor(0x1F, 0x29, 0x37)
COLOR_MUTED = RGBColor(0x6B, 0x72, 0x80)
HEX_SLATE_BG = "F1F5F9"
HEX_BORDER = "CBD5E1"

base_dir = "/Users/sameershekhar/Documents/GitHub/p1"
md_path = os.path.join(base_dir, "FINAL_COMPREHENSIVE_PROJECT_REPORT.md")
with open(md_path, "r", encoding="utf-8") as f:
    text = f.read()

lines = text.split("\n")
i = 0
table_lines = []

def process_table(tlines):
    if len(tlines) < 2:
        return
    headers = [c.strip() for c in tlines[0].strip().strip("|").split("|")]
    rows_data = []
    for r in tlines[2:]:
        if "|" in r:
            cells = [c.strip() for c in r.strip().strip("|").split("|")]
            rows_data.append(cells)
    
    t = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    
    # Header row
    hdr_cells = t.rows[0].cells
    for col_idx, htext in enumerate(headers):
        if col_idx < len(hdr_cells):
            clean_h = re.sub(r"[*_\$]", "", htext)
            hdr_cells[col_idx].text = clean_h
            for p in hdr_cells[col_idx].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.bold = True
                    run.font.size = Pt(9.5)
                    run.font.color.rgb = COLOR_SLATE
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_SLATE_BG}"/>')
            hdr_cells[col_idx]._tc.get_or_add_tcPr().append(shading)
    
    # Body rows
    for row_idx, rdata in enumerate(rows_data):
        row_cells = t.rows[row_idx + 1].cells
        for col_idx, ctext in enumerate(rdata):
            if col_idx < len(row_cells):
                clean_t = re.sub(r"[*_\$]", "", ctext)
                clean_t = clean_t.replace("&rarr;", "->").replace("&times;", "x")
                row_cells[col_idx].text = clean_t
                for p in row_cells[col_idx].paragraphs:
                    if col_idx > 0 and len(rdata) > 3:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.font.size = Pt(9)
                        run.font.color.rgb = COLOR_TEXT
    
    # Borders
    for row in t.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{HEX_BORDER}"/>
                    <w:left w:val="none"/>
                    <w:right w:val="none"/>
                </w:tcBorders>
            ''')
            tcPr.append(tcBorders)
    
    doc.add_paragraph()

while i < len(lines):
    line = lines[i]
    stripped = line.strip()
    
    # Tables
    if stripped.startswith("|") and "|" in stripped[1:]:
        table_lines.append(stripped)
        i += 1
        continue
    else:
        if table_lines:
            process_table(table_lines)
            table_lines = []
            
    # Images: ![caption](path)
    img_match = re.match(r"^!\[(.*?)\]\((.*?)\)$", stripped)
    if img_match:
        caption, img_rel_path = img_match.groups()
        img_abs_path = os.path.join(base_dir, img_rel_path) if not os.path.isabs(img_rel_path) else img_rel_path
        if os.path.exists(img_abs_path) and img_abs_path.endswith(('.png', '.jpg', '.jpeg')):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(2)
            run_img = p_img.add_run()
            run_img.add_picture(img_abs_path, width=Inches(6.0))
            if caption:
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(f"Figure: {caption}")
                run_cap.font.italic = True
                run_cap.font.size = Pt(9)
                run_cap.font.color.rgb = COLOR_MUTED
        i += 1
        continue

    # Headings
    if stripped.startswith("# "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(re.sub(r"[*_\$#]", "", stripped))
        run.font.bold = True
        run.font.size = Pt(18)
        run.font.color.rgb = COLOR_SLATE
    elif stripped.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(re.sub(r"[*_\$#]", "", stripped))
        run.font.bold = True
        run.font.size = Pt(14)
        run.font.color.rgb = COLOR_SLATE
    elif stripped.startswith("### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(re.sub(r"[*_\$#]", "", stripped))
        run.font.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = COLOR_SLATE
    elif stripped.startswith("#### "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(re.sub(r"[*_\$#]", "", stripped))
        run.font.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = COLOR_SLATE
    elif stripped.startswith("```"):
        code_lines = []
        i += 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            code_lines.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run("\n".join(code_lines))
        run.font.name = "Courier New"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
    elif stripped.startswith("* ") or stripped.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        clean_p = re.sub(r"^\s*[\*\-]\s*", "", line)
        clean_p = re.sub(r"[\$]", "", clean_p)
        run = p.add_run(sanitize(clean_p))
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT
    elif re.match(r"^\d+\.\s", stripped):
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(2)
        clean_p = re.sub(r"^\s*\d+\.\s*", "", line)
        clean_p = re.sub(r"[\$]", "", clean_p)
        run = p.add_run(sanitize(clean_p))
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT
    elif stripped.startswith("$$") and stripped.endswith("$$") and len(stripped) > 4:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        eq_text = stripped[2:-2].strip()
        run = p.add_run(sanitize(eq_text))
        run.font.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_SLATE
    elif stripped == "---" or stripped == "":
        pass
    else:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        clean_p = re.sub(r"[\$]", "", stripped)
        run = p.add_run(sanitize(clean_p))
        run.font.size = Pt(10)
        run.font.color.rgb = COLOR_TEXT

    i += 1

if table_lines:
    process_table(table_lines)

out_docx = os.path.join(base_dir, "FINAL_COMPREHENSIVE_PROJECT_REPORT.docx")
doc.save(out_docx)
print("SUCCESS: Document written to", out_docx)
